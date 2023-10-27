import os

import openpyxl
from docx import Document

import sys
sys.path.append(os.getcwd()+'/src/utils')
import excel_utils
import word_utils


# Load EXCEL
excel_file_path = "public/excel.xlsx"
file = openpyxl.load_workbook(excel_file_path)
sheet = file.active


# Get cell value
cell_A1_value = excel_utils.get_cell_value(sheet, 'A1')


# Create WORD
doc = Document()


# Add content
word_utils.add_paragraph(doc, f"This is {cell_A1_value}")


# Save WORD
word_file_path = "public/word.docx"
doc.save(word_file_path)

print(f"Word document saved to {word_file_path}")